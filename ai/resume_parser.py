"""Resume parser — extracts plain text from PDF and DOCX files."""

import io

from utils.text_cleaner import clean_text


def parse_resume(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a PDF or DOCX resume file.

    Args:
        file_bytes: raw bytes of the uploaded file
        filename: original filename (used to determine parser)

    Returns:
        Cleaned plain text string
    """
    fname = filename.lower()
    if fname.endswith(".pdf"):
        return _parse_pdf(file_bytes)
    elif fname.endswith(".docx"):
        return _parse_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file format: {filename}. Please upload PDF or DOCX.")


def _parse_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
    except ImportError:
        raise ImportError("pdfplumber is required for PDF parsing. Install it with: pip install pdfplumber")

    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

    raw = "\n".join(text_parts)
    return clean_text(raw)


def _parse_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing. Install it with: pip install python-docx")

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    raw = "\n".join(paragraphs)
    return clean_text(raw)
